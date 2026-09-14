"""Generate the editable DOCX Declarations document (Array submission)."""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Declarations")
r.bold = True
r.font.size = Pt(15)

doc.add_paragraph()


def heading(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    return p


def body(text):
    doc.add_paragraph(text)


heading("Funding")
body("This research received no specific grant from any funding agency in "
     "the public, commercial, or not-for-profit sectors.")

heading("Conflict of interest")
body("The authors declare no competing interests. A separate Declaration of "
     "Competing Interests has been submitted alongside this document.")

heading("Data and code availability")
body("Code and data to reproduce all results are available at "
     "https://github.com/vinhqdang/KAN_time_series.")

heading("Declaration of generative AI and AI-assisted technologies in the "
        "writing process")
body("During the preparation of this work, the authors used an AI-based "
     "writing assistant to support the drafting process. After using this "
     "tool, the authors reviewed and edited the content as needed and take "
     "full responsibility for the content of the publication.")

doc.add_paragraph()
doc.add_paragraph("Quang-Vinh Dang, Dat Le, Minh Ngoc Dinh.")

doc.save("declarations.docx")
print("Saved declarations.docx")
